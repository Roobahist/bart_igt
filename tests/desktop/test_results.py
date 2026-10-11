import copy
import importlib
import json
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from study import settings as config
from apps.desktop.results import SessionResults, prepare_config, recover, task_order


class SessionTests(unittest.TestCase):
    def test_task_order(self):
        for value in ("2", "", "participant", None):
            self.assertEqual(task_order(value), ("IGT", "BART"))
        self.assertEqual(task_order("3"), ("BART", "IGT"))

    def test_import_does_not_require_gui(self):
        from apps.desktop import experiment as main

        importlib.reload(main)

    def test_assets_resolve_outside_repository(self):
        cfg = SimpleNamespace(
            **{
                k: copy.deepcopy(v)
                for k, v in vars(config).items()
                if not k.startswith("_") and k != "sys"
            }
        )
        with tempfile.TemporaryDirectory() as directory:
            original = Path.cwd()
            try:
                os.chdir(directory)
                prepare_config(cfg)
            finally:
                os.chdir(original)
        self.assertTrue(Path(cfg.bart_balloon_path).is_file())
        self.assertTrue(Path(cfg.video_path).is_file())

    def test_invalid_settings(self):
        for name, value in [
            ("bart_max_balloon_capacity", 0),
            ("bart_blowing_factor", 0),
            ("igt_total_number_of_trials", -1),
            ("bart_blow_addition", float("nan")),
            ("column_order", ["trial"]),
        ]:
            cfg = SimpleNamespace(
                **{k: v for k, v in vars(config).items() if not k.startswith("_")}
            )
            setattr(cfg, name, value)
            with self.subTest(name=name), self.assertRaises(ValueError):
                prepare_config(cfg)

    def test_excel_roundtrip_and_literal_participant_text(self):
        with tempfile.TemporaryDirectory() as directory:
            result = SessionResults(directory, config.column_order[::-1])
            result.append({"subjectCode": "../../escape", "subjectName": "=1+1"})
            result.append({"section": "BART", "trial": 1, "gain": 0.05, "pumps": 1})
            path = result.save()
            result.close()
            self.assertEqual(path.parent, Path(directory))
            workbook = load_workbook(path)
            sheet = workbook.active
            self.assertEqual([c.value for c in sheet[1]], config.column_order[::-1])
            name = next(c for c in sheet[2] if c.value == "=1+1")
            self.assertEqual(name.data_type, "s")
            workbook.close()
            lines = result.journal_path.read_text().splitlines()
            self.assertEqual(json.loads(lines[2])["gain"], 0.05)

    def test_failed_export_retains_journal_and_existing_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            result = SessionResults(directory, ["trial"])
            result.append({"trial": 1})
            result.save()
            original = result.path.read_bytes()
            result.append({"trial": 2})
            with patch("openpyxl.Workbook.save", side_effect=OSError("disk full")):
                with self.assertRaises(OSError):
                    result.save()
            result.close()
            self.assertEqual(result.path.read_bytes(), original)
            recovered = load_workbook(recover(result.journal_path))
            self.assertEqual(recovered.active.max_row, 3)
            recovered.close()
            self.assertEqual(len(list(Path(directory).glob("*.xlsx"))), 2)

    def test_recovery_only_tolerates_truncated_final_record(self):
        with tempfile.TemporaryDirectory() as directory:
            result = SessionResults(directory, ["trial"])
            result.append({"trial": 1})
            result.close()
            with result.journal_path.open("a") as stream:
                stream.write('{"trial":')
            recovered = load_workbook(recover(result.journal_path))
            self.assertEqual(recovered.active.max_row, 2)
            recovered.close()
            with result.journal_path.open("a") as stream:
                stream.write("\n")
            with self.assertRaises(json.JSONDecodeError):
                recover(result.journal_path)

    def test_unknown_fields_do_not_write_a_response(self):
        with tempfile.TemporaryDirectory() as directory:
            result = SessionResults(directory, ["trial"])
            try:
                with self.assertRaises(ValueError):
                    result.append({"typo": 1})
                self.assertEqual(len(result.journal_path.read_text().splitlines()), 1)
            finally:
                result.close()
